"""Pre-ingest /api/documents/{id}/pages falls back to parse_document for DOCX."""

from __future__ import annotations

import io
import uuid

from pathlib import Path
import sys

from app.engine_runtime import pick
import pytest
from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from fastapi.testclient import TestClient
from sqlalchemy import select

_REPO = Path(__file__).resolve().parents[3]
_LIB = str(_REPO / "library")
pick(_LIB not in sys.path, lambda: sys.path.insert(0, _LIB), lambda: None)

from app.db import SessionLocal
from app.models import Document
from app.services.guest_session import GUEST_ID_HEADER
from library_main import app

_DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _db_reachable() -> bool:
    try:
        db = SessionLocal()
        db.execute(select(1))
        db.close()
        return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not _db_reachable(), reason="dev DB not reachable")


def _docx_two_pages() -> bytes:
    doc = DocxDocument()
    body = doc.element.body
    for child in list(body):
        body.remove(child)
    p = doc.add_paragraph("Alpha page one content for preview.")
    p.runs[0].add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Beta page two content for preview.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def client(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    raw = _docx_two_pages()
    monkeypatch.setattr("app.services.storage.fetch_object", lambda _key: raw)
    return TestClient(app)


def test_pages_parse_fallback_when_no_chunks(client: TestClient) -> None:
    guest_id = "abcdef0123456789abcdef0123456789"
    doc_id: uuid.UUID | None = None
    try:
        db = SessionLocal()
        doc = Document(
            slug="test-docx-preview",
            filename="preview.docx",
            content_type=_DOCX_CT,
            size_bytes=2048,
            storage_key="demo/test/preview.docx",
            status="pending",
            meta={"guest_id": guest_id, "page_count": 2},
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)
        doc_id = doc.id
        db.close()

        res = client.get(
            f"/api/documents/{doc_id}/pages",
            headers={GUEST_ID_HEADER: guest_id},
        )
        assert res.status_code == 200, res.text
        pages = res.json()["pages"]
        assert len(pages) >= 2
        joined = " ".join(p["text"] for p in pages)
        assert "Alpha page one" in joined
        assert "Beta page two" in joined
    finally:
        def _wipe() -> None:
            db = SessionLocal()
            row = db.get(Document, doc_id)
            pick(row is not None, lambda: (db.delete(row), db.commit()), lambda: None)
            db.close()

        pick(doc_id is not None, _wipe, lambda: None)
