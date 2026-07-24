"""DOCX / text parsing must produce multi-page study units like PDF / paste."""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK

from app.services.parse import count_document_pages, parse_document, parse_document_page
from app.services.web_import import READER_PAGE_CHARS


def _docx_bytes(*builders: object) -> bytes:
    doc = DocxDocument()
    # Drop the default empty paragraph so builders own the body.
    body = doc.element.body
    for child in list(body):
        body.remove(child)
    for build in builders:
        build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_paragraph(text: str):
    def build(doc: DocxDocument) -> None:
        doc.add_paragraph(text)

    return build


def _add_page_break_then(text: str):
    def build(doc: DocxDocument) -> None:
        p = doc.add_paragraph("before")
        p.runs[0].add_break(WD_BREAK.PAGE)
        doc.add_paragraph(text)

    return build


def test_parse_docx_splits_on_hard_page_break() -> None:
    data = _docx_bytes(
        _add_paragraph("First page body"),
        _add_page_break_then("Second page body"),
    )
    pages = parse_document(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data,
    )
    assert len(pages) >= 2
    assert pages[0]["page"] == 1
    assert "First page" in pages[0]["text"]
    assert any("Second page" in p["text"] for p in pages)


def test_parse_docx_soft_paginates_long_doc_without_breaks() -> None:
    long = "\n\n".join(f"Section {i}. " + ("word " * 80) for i in range(12))
    data = _docx_bytes(_add_paragraph(long))
    pages = parse_document(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data,
    )
    assert len(pages) > 1
    assert all(pages[i]["page"] == i + 1 for i in range(len(pages)))
    joined = "\n\n".join(p["text"] for p in pages)
    assert "Section 0." in joined
    assert "Section 11." in joined
    assert all(len(p["text"]) <= READER_PAGE_CHARS + 200 for p in pages)


def test_count_document_pages_matches_docx_parse() -> None:
    long = "\n\n".join(f"Block {i}. " + ("x " * 100) for i in range(10))
    data = _docx_bytes(_add_paragraph(long))
    ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert count_document_pages(ct, data) == len(parse_document(ct, data))


def test_parse_document_page_picks_docx_page() -> None:
    data = _docx_bytes(
        _add_paragraph("Alpha page"),
        _add_page_break_then("Beta page"),
    )
    ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    page2 = parse_document_page(ct, data, 2)
    assert page2["page"] == 2
    assert "Beta" in page2["text"]


def test_parse_plain_text_soft_paginates() -> None:
    text = "\n\n".join(f"Para {i}. " + ("word " * 100) for i in range(10))
    pages = parse_document("text/plain", text.encode("utf-8"))
    assert len(pages) > 1


def test_avidpay_like_docx_is_multi_page() -> None:
    """Regression: long Word KT notes must not collapse to a single study page."""
    from pathlib import Path

    path = Path("/Users/benito/Desktop/AvidPay KT - Complete Reference.docx")
    if not path.is_file():
        return
    data = path.read_bytes()
    ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    pages = parse_document(ct, data)
    assert len(pages) >= 8
    assert all(pages[i]["page"] == i + 1 for i in range(len(pages)))
    assert any("Part 1" in p["text"] for p in pages)
