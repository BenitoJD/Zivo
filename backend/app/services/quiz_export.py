"""Render a generated quiz to a .docx worksheet (student or teacher/answer-key version)."""

from __future__ import annotations

import io

from app.engine_runtime import Pred, Rule, apply, choose, first_match, pick

_LETTERS = "ABCDEFGH"

# Single-best-answer MCQ variants share the plain "mcq" payload (options +
# answer_index), so they render and export through the same branch.
_SINGLE_ANSWER_MCQ = ("mcq", "mcq_negative", "assertion_reason", "scenario", "cloze")

_ANSWER_RULES = (
    Rule(when=(Pred("single", "truthy"),), action="single"),
    Rule(when=(Pred("type", "eq", "multi"),), action="multi"),
    Rule(when=(Pred("type", "eq", "truefalse"),), action="truefalse"),
    Rule(when=(Pred("fill", "truthy"),), action="fill"),
    Rule(when=(Pred("type", "eq", "matching"),), action="matching"),
    Rule(when=(), action="empty"),
)
_BODY_RULES = (
    Rule(when=(Pred("options", "truthy"),), action="options"),
    Rule(when=(Pred("type", "eq", "truefalse"),), action="truefalse"),
    Rule(when=(Pred("type", "eq", "matching"),), action="matching"),
    Rule(when=(Pred("fill_short", "truthy"),), action="fill_short"),
    Rule(when=(Pred("type", "eq", "essay"),), action="essay"),
    Rule(when=(), action="none"),
)


def _answer_text(q: dict) -> str:
    t = q.get("type")
    hit = first_match(
        _ANSWER_RULES,
        {
            "type": t,
            "single": t in _SINGLE_ANSWER_MCQ,
            "fill": t in ("fill_blank", "short", "essay"),
        },
    )

    def _single() -> str:
        i = q.get("answer_index", 0)
        opts = q.get("options", [])
        return pick(i < len(opts), lambda: f"{_LETTERS[i]}. {opts[i]}", lambda: _LETTERS[i])

    def _multi() -> str:
        idxs = q.get("answer_indices", [])
        opts = q.get("options", [])
        return ", ".join(
            f"{_LETTERS[i]}. {opts[i]}" for i in filter(lambda i: i < len(opts), idxs)
        )

    return apply(
        hit.action,
        {
            "single": _single,
            "multi": _multi,
            "truefalse": lambda: choose(bool(q.get("answer")), "True", "False"),
            "fill": lambda: str(q.get("answer") or ""),
            "matching": lambda: "; ".join(
                f"{p.get('left')} → {p.get('right')}" for p in q.get("pairs", [])
            ),
            "empty": lambda: "",
        },
    )


def build_quiz_docx(*, title: str, questions: list[dict], with_answers: bool) -> bytes:
    """Build a worksheet docx. with_answers=True produces the teacher/answer-key version."""
    from docx import Document as Docx
    from docx.shared import Pt

    doc = Docx()
    doc.add_heading(title or "Quiz", level=0)
    sub = doc.add_paragraph(choose(with_answers, "Answer key", "Name: ____________________    Date: __________"))
    sub.runs[0].italic = True

    for n, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"{n}. {q.get('prompt', '')}")
        run.bold = True
        run.font.size = Pt(11)

        t = q.get("type")
        hit = first_match(
            _BODY_RULES,
            {
                "type": t,
                "options": t in _SINGLE_ANSWER_MCQ or t == "multi",
                "fill_short": t in ("fill_blank", "short"),
            },
        )

        def _options() -> None:
            for i, opt in enumerate(q.get("options", [])):
                doc.add_paragraph(f"{_LETTERS[i]}. {opt}", style="List Bullet")

        def _matching() -> None:
            lefts = [p_.get("left", "") for p_ in q.get("pairs", [])]
            rights = [p_.get("right", "") for p_ in q.get("pairs", [])]
            for i, left in enumerate(lefts):
                doc.add_paragraph(f"{i + 1}. {left}    ____", style="List Bullet")
            for j, right in enumerate(rights):
                doc.add_paragraph(f"   {_LETTERS[j]}. {right}")

        def _essay() -> None:
            for _ in range(4):
                doc.add_paragraph("__________________________________________")

        apply(
            hit.action,
            {
                "options": _options,
                "truefalse": lambda: doc.add_paragraph("True  /  False"),
                "matching": _matching,
                "fill_short": lambda: doc.add_paragraph("__________________________________________"),
                "essay": _essay,
                "none": lambda: None,
            },
        )

        def _answers() -> None:
            ap = doc.add_paragraph()
            ar = ap.add_run(f"Answer: {_answer_text(q)}")
            ar.italic = True
            expl = (q.get("explanation") or "").strip()

            def _expl() -> None:
                ep = doc.add_paragraph()
                er = ep.add_run(expl)
                er.italic = True
                er.font.size = Pt(9)

            pick(bool(expl), _expl, lambda: None)

        pick(with_answers, _answers, lambda: None)
        doc.add_paragraph("")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
